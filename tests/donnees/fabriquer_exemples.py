"""Fabrique les fichiers d'exemple (clients fictifs) utilisés par les tests d'import.

    python tests/donnees/fabriquer_exemples.py

Produit :
- clients-exemple.xlsx : ancienne présentation (Client | Adresse | Contact | Logo),
  avec un logo et une rubrique « Techniciens » ;
- clients-contacts-exemple.xlsx : présentation actuelle, une ligne par contact
  (Client | Titre | Nom | Prénom | Fonction | Téléphone | Mobile | Email), avec une
  cellule Client fusionnée sur deux contacts et un mobile saisi comme un nombre ;
- clients-exemple-excel.csv : CSV « ; » en encodage Windows, comme l'enregistre Excel. Les fichiers produits sont versionnés :
les tests n'ont pas besoin d'openpyxl ni de Pillow.
"""
import io
import os

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ImageXL
from PIL import Image, ImageDraw

ICI = os.path.dirname(os.path.abspath(__file__))

CLIENTS = [
    ('FROMAGERIE_DU_JURA', '12 rue des Sapins\n39000 LONS-LE-SAUNIER',
     'Paul MARTIN - Responsable maintenance 06 12 34 56 78'),
    ('LAITERIE_DES_ALPES', '5 avenue du Lac\n74000 ANNECY', 'Julie BERNARD'),
    ('BISCUITERIE_NORD', '8 rue du Port\n59000 LILLE', None),
    ('CONSERVERIE_ÉTÉ', '3 chemin des Prés\n01150 BLYES', 'Hélène FAURE - Chef d\'atelier'),
]


def logo_png():
    im = Image.new('RGB', (120, 60), (11, 61, 145))
    d = ImageDraw.Draw(im)
    d.rectangle([8, 8, 112, 52], outline=(255, 255, 255), width=3)
    d.text((30, 22), 'LOGO', fill=(255, 255, 255))
    b = io.BytesIO()
    im.save(b, 'PNG')
    b.seek(0)
    return b


def xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Feuille 1'
    ws.append(['COORDONNÉES CLIENTS'])
    ws.append(['Client', 'Adresse', 'Contact', 'Logo'])
    for c in CLIENTS:
        ws.append(list(c))
    ws.add_image(ImageXL(logo_png()), 'D3')        # logo de FROMAGERIE_DU_JURA
    ws.append([])
    ws.append(['Techniciens'])
    ws.append(['Jean DUPONT'])
    ws.append(['Marie DURAND'])
    ws.append([])
    ws.append([])
    ws.append(["Nuits d'hotel"])
    ws.append([1])
    ws.append([2])
    wb.save(os.path.join(ICI, 'clients-exemple.xlsx'))


def xlsx_contacts():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Contacts'
    ws.append(['COORDONNÉES CLIENTS'])
    ws.append(['Client', 'Titre', 'Nom', 'Prénom', 'Fonction', 'Téléphone', 'Mobile', 'Email'])
    ws.append(['FROMAGERIE_DU_JURA', 'M.', 'MARTIN', 'Paul', 'Responsable maintenance',
               '03 84 00 00 00', 612345678, 'paul.martin@exemple.fr'])     # mobile saisi comme un nombre
    ws.append([None, 'Mme', 'LEROY', 'Sophie', 'Responsable production',
               None, '06.98.76.54.32', 's.leroy@exemple.fr'])
    ws.merge_cells('A3:A4')                                                  # même client, deux contacts
    ws.append(['LAITERIE_DES_ALPES', 'Mme', 'BERNARD', 'Julie', 'Directrice technique',
               '04 50 00 00 00', None, 'j.bernard@exemple.fr'])
    ws.append(['NOUVEAU_CLIENT_SA', 'M.', 'PETIT', 'Luc', 'Acheteur', None, '0611223344', 'luc.petit@exemple.fr'])
    wb.save(os.path.join(ICI, 'clients-contacts-exemple.xlsx'))


def csv_excel():
    lignes = ['COORDONNÉES CLIENTS;;;', 'Client;Adresse;Contact;Logo']
    for nom, adr, contact in CLIENTS:
        lignes.append('%s;"%s";%s;' % (nom, adr, contact or ''))
    with open(os.path.join(ICI, 'clients-exemple-excel.csv'), 'wb') as f:
        f.write('\r\n'.join(lignes).encode('cp1252'))


if __name__ == '__main__':
    xlsx()
    xlsx_contacts()
    csv_excel()
    print('fichiers d\'exemple écrits dans', ICI)
